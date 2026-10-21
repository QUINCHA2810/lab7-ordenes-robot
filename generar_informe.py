"""
Laboratorio 7 - Genera informe.docx a partir de tablas.json (que produce tablas.py) y de los
archivos de resultados. Ninguna cifra esta escrita a mano: todas salen de esos archivos.

    py -3.12 tablas.py --json tablas.json
    py -3.12 generar_informe.py
"""
import json

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import textos_informe as T

with open("tablas.json", encoding="utf-8") as f:
    TB = json.load(f)


def fmt(v, d=3, pct=False):
    if v is None:
        return "n/d"
    if isinstance(v, (int, float)):
        return f"{v:.1f}%" if pct else f"{v:.{d}f}"
    return str(v)


def configurar(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.59), Cm(27.94)
    s.left_margin = s.right_margin = Cm(1.8)
    s.top_margin = s.bottom_margin = Cm(1.6)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(9.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing = 1.0
    for nombre, tam in (("Heading 1", 11.5), ("Heading 2", 10)):
        h = doc.styles[nombre]
        h.font.size = Pt(tam)
        h.font.name = "Calibri"
        h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        h.paragraph_format.space_before = Pt(6)
        h.paragraph_format.space_after = Pt(2)


def parrafo(doc, texto, negrita_inicio=None, cursiva=False, tam=None, alinear=True):
    p = doc.add_paragraph()
    if negrita_inicio:
        r = p.add_run(negrita_inicio)
        r.bold = True
        if tam:
            r.font.size = Pt(tam)
    r = p.add_run(texto)
    r.italic = cursiva
    if tam:
        r.font.size = Pt(tam)
    if alinear:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def sombrear(celda, color="D9E2F3"):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def tabla(doc, cabecera, filas, titulo, ancho=None, resaltar=None):
    cap = doc.add_paragraph()
    r = cap.add_run(titulo)
    r.bold = True
    r.font.size = Pt(8.5)
    cap.paragraph_format.space_after = Pt(1)
    t = doc.add_table(rows=1, cols=len(cabecera))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cabecera):
        cel = t.rows[0].cells[i]
        cel.text = ""
        run = cel.paragraphs[0].add_run(c)
        run.bold = True
        run.font.size = Pt(8)
        sombrear(cel)
    for k, fila in enumerate(filas):
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = ""
            run = celdas[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(8)
            if resaltar is not None and k == resaltar:
                run.bold = True
    if ancho:
        for fila in t.rows:
            for i, w in enumerate(ancho):
                fila.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def main():
    doc = Document()
    configurar(doc)

    tit = doc.add_paragraph()
    tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tit.add_run("Comprensión de órdenes habladas en español para un robot de servicio con LLMs ligeros")
    r.bold = True
    r.font.size = Pt(13)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Laboratorio 7 — NLP y LLMs · Universidad Santo Tomás\n"
                    "[INTEGRANTE 1] · [INTEGRANTE 2] · [INTEGRANTE 3] · Fecha: [FECHA DE ENTREGA]")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    for seccion in T.SECCIONES(TB, fmt):
        tipo = seccion[0]
        if tipo == "h1":
            doc.add_heading(seccion[1], level=1)
        elif tipo == "h2":
            doc.add_heading(seccion[1], level=2)
        elif tipo == "p":
            parrafo(doc, seccion[1])
        elif tipo == "pb":
            parrafo(doc, seccion[2], negrita_inicio=seccion[1])
        elif tipo == "tabla":
            tabla(doc, *seccion[1:])
        elif tipo == "ref":
            p = parrafo(doc, seccion[1], tam=8.5, alinear=False)
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(1)
    doc.save("informe.docx")
    print("informe.docx generado")


if __name__ == "__main__":
    main()
